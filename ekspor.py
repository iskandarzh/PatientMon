"""Pembuatan file export hasil pengukuran (Excel & PDF)."""
from io import BytesIO
from xml.sax.saxutils import escape

KOLOM = [
    ("Waktu", 17), ("No. RM", 10), ("Nama Pasien", 24), ("Sistolik", 10), ("Diastolik", 11), ("Nadi", 7),
    ("Kategori Tensi", 19), ("Gula Darah", 11), ("Jenis Gula", 20), ("Kategori Gula", 15), ("Catatan", 40),
]


def baris_pengukuran(m, kategori_tensi, kategori_gula, jenis_gula):
    kt = kategori_tensi(m["sistolik"], m["diastolik"])
    kg = kategori_gula(m["gula_darah"], m["jenis_gula"])
    return [
        m["waktu"].replace("T", " "), m["no_rm"] or "", m["nama_pasien"],
        m["sistolik"], m["diastolik"], m["nadi"], kt[0] if kt else "",
        m["gula_darah"], jenis_gula.get(m["jenis_gula"], "") if m["gula_darah"] is not None else "",
        kg[0] if kg else "", m["catatan"] or "",
    ]


def buat_xlsx(rows, judul, keterangan):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Pengukuran"
    ws.append([judul])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([keterangan])
    ws.append([])
    ws.append([k for k, _ in KOLOM])
    header = ws.max_row
    for cell in ws[header]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D6EFD")
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                # Teks berawalan "=" jangan sampai dianggap rumus oleh Excel.
                cell.data_type = "s"
    for i, (_, lebar) in enumerate(KOLOM, start=1):
        ws.column_dimensions[ws.cell(row=header, column=i).column_letter].width = lebar
    ws.freeze_panes = ws.cell(row=header + 1, column=1)
    ws.auto_filter.ref = f"A{header}:{ws.cell(row=header, column=len(KOLOM)).column_letter}{max(ws.max_row, header)}"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def buat_pdf(rows, judul, keterangan):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    sel = styles["BodyText"].clone("sel", fontSize=7.5, leading=9)
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), title=judul,
        leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
    )

    def p(v):
        return Paragraph(escape("" if v is None else str(v)), sel)

    data = [[p(k) for k, _ in KOLOM]] + [[p(v) for v in row] for row in rows]
    if not rows:
        data.append([p("Tidak ada data")] + [""] * (len(KOLOM) - 1))
    total = sum(w for _, w in KOLOM)
    lebar = [doc.width * w / total for _, w in KOLOM]
    tabel = Table(data, colWidths=lebar, repeatRows=1)
    tabel.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D6EFD")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FA")]),
    ]))
    doc.build([
        Paragraph(escape(judul), styles["Title"]),
        Paragraph(escape(keterangan), styles["Normal"]),
        Spacer(1, 6 * mm),
        tabel,
    ])
    buf.seek(0)
    return buf
